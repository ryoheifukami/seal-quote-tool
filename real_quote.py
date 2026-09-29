# -*- coding: utf-8 -*-
"""
見積作成（実際のExcel見積フォーマットを再現）— Streamlit UI

pricing_real.py の検証済みエンジンで、平圧・間欠／マスキング／コニカミノルタの3方式を計算し、
宛先・件名等を入れて見積書（テキスト／Excel）まで出力する。app.py の「見積作成」から render()。
"""

from io import BytesIO
from datetime import date, datetime, timedelta

import streamlit as st

import os

import pricing_real as pr
import pdf_intake
import intake_parse
import store
import masters_ui
import machine_select
from quote_excel import ISSUER


def _get_api_key():
    try:
        k = st.secrets.get("ANTHROPIC_API_KEY", "")
    except Exception:
        k = ""
    return k or os.environ.get("ANTHROPIC_API_KEY", "")


def _paste_to_real(fields):
    """貼り付け文の抽出結果 → 実計算の入力キーへ変換。(method, 更新dict, 不足項目リスト) を返す。"""
    is_tape = fields.get("product") in ("マスキングテープ", "オリジナル紙テープ") or fields.get("tape_width_mm")
    if is_tape:
        method = "マスキングテープ"
        upd = {}
        w = fields.get("tape_width_mm") or fields.get("width_mm")
        if w:
            upd["m_w"] = float(w)
        if fields.get("tape_length_m"):
            upd["m_mk"] = float(fields["tape_length_m"])
        if fields.get("qty"):
            upd["m_qty"] = int(fields["qty"])
        missing = []
        if "m_w" not in upd:
            missing.append("製品の幅（mm）")
        if "m_mk" not in upd:
            missing.append("巻きメーター（m）")
        if "m_qty" not in upd:
            missing.append("数量（巻）")
    else:
        method = "平圧・間欠印刷"
        upd = {}
        if fields.get("width_mm"):
            upd["h_w"] = float(fields["width_mm"])
        if fields.get("height_mm"):
            upd["h_h"] = float(fields["height_mm"])
        if fields.get("qty"):
            upd["h_qty"] = int(fields["qty"])
        col = None
        if fields.get("color_type") == "単色":
            col = int(fields.get("spot_colors") or 1)
        elif fields.get("color_type") == "フルカラー":
            col = 4
        if col:
            upd["h_col"] = min(5, max(1, col))
        missing = []
        if "h_w" not in upd:
            missing.append("寸法ヨコ W（mm）")
        if "h_h" not in upd:
            missing.append("寸法タテ P（mm）")
        if "h_qty" not in upd:
            missing.append("数量（枚）")
        if "h_col" not in upd:
            missing.append("色数（何色か／フルカラーか単色か）")
    return method, upd, missing


def render_paste_flow():
    """問い合わせ文を貼り付け → 抽出 → 仕様に反映＋不足項目の確認表示。"""
    st.markdown("#### 📋 問い合わせ文を貼り付け")
    st.caption("メールやフォームの問い合わせ文を貼り付けて「読み取り」を押すと、下の仕様に自動反映します。"
               "必要な項目が足りない場合は、確認すべき点を表示します。")
    st.text_area("問い合わせ文をそのまま貼り付け", key="rq_paste", height=150,
                 placeholder="例）\nお世話になります。下記で見積りをお願いします。\n"
                             "サイズ：50×30mm／数量：3000枚／単色1色（黒）／透明PET")
    if st.button("読み取り → 仕様に反映", type="primary", key="rq_paste_go"):
        fields, method_used, notes = intake_parse.parse_inquiry(st.session_state.get("rq_paste", ""), _get_api_key())
        method, upd, missing = _paste_to_real(fields)
        st.session_state["rq_method"] = method
        for k, v in upd.items():
            st.session_state[k] = v
        st.session_state["rq_paste_missing"] = missing
        st.session_state["rq_paste_method_read"] = method_used
        st.session_state["rq_paste_done"] = True
        st.rerun()
    if st.session_state.get("rq_paste_done"):
        st.caption(f"読み取り方式：{st.session_state.get('rq_paste_method_read','')}")
        missing = st.session_state.get("rq_paste_missing", [])
        if missing:
            st.warning("⚠ 次の項目が読み取れませんでした。**お客様にご確認ください**：\n\n"
                       + "\n".join(f"- {m}" for m in missing)
                       + "\n\n（下の仕様欄に直接入力しても計算できます）")
        else:
            st.success("✅ 見積りに必要な項目がそろいました。下の内容で金額が出ます。")

METHOD_PREFIX = {"平圧・間欠印刷": "h_", "マスキングテープ": "m_", "コニカミノルタ オンデマンド": "k_"}


def _snapshot(method):
    pre = METHOD_PREFIX.get(method, "")
    snap = {k: v for k, v in st.session_state.items()
            if isinstance(k, str) and k.startswith(pre)}
    snap["rq_name"] = st.session_state.get("rq_name", "")
    return snap


def _apply_template(t):
    for k, v in (t.get("inputs") or {}).items():
        st.session_state[k] = v
    if t.get("method"):
        st.session_state["rq_method"] = t["method"]


def _template_bar():
    ts = store.get_templates()
    with st.expander("⭐ テンプレート（入力プリセット）", expanded=False):
        if ts:
            c = st.columns([3, 1])
            sel = c[0].selectbox("保存済みテンプレート", [t["name"] for t in ts], key="tpl_sel")
            if c[1].button("読込", key="tpl_load"):
                _apply_template(next(x for x in ts if x["name"] == sel))
                st.rerun()
        else:
            st.caption("まだテンプレートがありません。下の欄で現在の入力を保存できます。")
        c2 = st.columns([3, 1])
        nm = c2[0].text_input("現在の入力をテンプレート保存（名前）", key="tpl_name")
        if c2[1].button("保存", key="tpl_save"):
            if nm:
                method = st.session_state.get("rq_method", "平圧・間欠印刷")
                store.save_template(nm, method, _snapshot(method))
                st.success(f"テンプレート「{nm}」を保存しました。")

TAX_RATE = 0.10


def yen(v):
    try:
        return f"¥{round(v):,}"
    except Exception:
        return "—"


def _sel(label, mapping, key, help=None):
    chosen = st.selectbox(label, list(mapping.keys()), key=key, help=help)
    return mapping[chosen]


def _breakdown_table(breakdown):
    rows = [{"項目": lbl, "原価": round(v)} for lbl, v in breakdown if v]
    if rows:
        st.dataframe(rows, width="stretch", hide_index=True,
                     column_config={"原価": st.column_config.NumberColumn(format="¥%d")})


def _machine_block(candidates, excluded, haku_candidates, restrict_category=None):
    """印刷機選定の候補を表示する共通ブロック。restrict_categoryを指定すると、その方式の機種のみに絞る。"""
    if restrict_category:
        candidates = [m for m in candidates if m["category"] == restrict_category]
        excluded = [m for m in excluded if m["category"] == restrict_category]
    with st.expander("🖨️ 印刷機選定の候補（現在の仕様から自動判定・参考）"):
        if candidates:
            for m in candidates:
                st.markdown(f"**✅ {m['name']}**（{m['category']}）")
                st.caption(
                    f"{m['desc']}／色数対応：{m['color_mode']}"
                    + (f"（{m['max_colors']}色まで）" if m['max_colors'] else "")
                    + (f"／NG色：{m['ng_color']}" if m.get('ng_color') else "")
                )
                st.caption(f"面付目安：{m['menzuke_note']}")
        else:
            st.warning("現在の仕様に完全に合う印刷機は見つかりませんでした。下の対象外理由をご確認ください。")
        if haku_candidates:
            st.markdown("**🔶 箔押し工程の候補**")
            for m in haku_candidates:
                st.markdown(f"- {m['name']}（最大 H{m['max_h']}×W{m['max_w']}mm）")
        if excluded:
            with st.expander("対象外の機種と理由"):
                for m in excluded:
                    st.caption(f"✕ {m['name']}：" + "／".join(m["exclude_reasons"]))
        st.caption("※ 最終的な機種判断は担当者の方でご確認ください。")


def _cost_block(r):
    c = st.columns(4)
    c[0].metric("原価合計", yen(r["cost_total"]))
    c[1].metric("材料粗利後", yen(r["material_after"]))
    c[2].metric("その他粗利後", yen(r["other_after"]))
    c[3].metric("参考単価（誤差あり）", f"¥{r['ref_unit']:.2f}")
    with st.expander("原価の内訳を見る"):
        _breakdown_table(r["breakdown"])


# =====================================================================
# 各方式の入力・計算（返り値：result, spec_text, unit）
# =====================================================================
def render_hiraatsu():
    st.markdown("#### 基本仕様")
    a = st.columns(4)
    width = a[0].number_input("寸法ヨコ W (mm)", min_value=0.0, value=50.0, step=1.0, key="h_w")
    height = a[1].number_input("寸法タテ P (mm)", min_value=0.0, value=30.0, step=1.0, key="h_h")
    colors = a[2].number_input("色数 (1〜5)", min_value=1, max_value=5, value=1, step=1, key="h_col")
    machine = _sel("機種", {"平圧": 2, "間欠": 3}, "h_mac")
    b = st.columns(4)
    mw = b[0].number_input("面付ヨコ", min_value=1, value=1, step=1, key="h_mw")
    mp = b[1].number_input("面付タテ", min_value=1, value=1, step=1, key="h_mp")
    nuki = _sel("抜き", {"ハーフカット": 2, "全抜き": 1}, "h_nuki")
    finish = _sel("仕上形状", {"シート": 1, "ロール": 2}, "h_fin")

    st.markdown("#### 材料（平米単価）")
    masters_ui.material_picker("h_gt")
    m = st.columns(3)
    genshi_tanka = m[0].number_input("原紙 平米単価", min_value=0.0, step=1.0, key="h_gt")
    pp_tanka = m[1].number_input("PP 平米単価", min_value=0.0, value=0.0, step=1.0, key="h_pt")
    haku_tanka = m[2].number_input("箔 平米単価", min_value=0.0, value=0.0, step=1.0, key="h_ht")

    st.markdown("#### 加工・オプション")
    o = st.columns(4)
    numbering = 1 if o[0].checkbox("ナンバリング印字", key="h_num") else 0
    ura = 1 if o[1].checkbox("ウラ(セパ)印刷", key="h_ura") else 0
    half = 1 if o[2].checkbox("ハーフラミ加工", key="h_half") else 0
    emb = 1 if o[3].checkbox("エンボス加工", key="h_emb") else 0
    o2 = st.columns(3)
    plate_change = o2[0].number_input("版替回数", min_value=0, value=0, step=1, key="h_pc")
    color_change = o2[1].number_input("色替回数", min_value=0, value=0, step=1, key="h_cc")
    add_process = o2[2].number_input("追加工程数", min_value=0, value=0, step=1, key="h_ap")

    cands, excl, haku_cands = machine_select.suggest_machines(
        colors=int(colors), width=width, height=height,
        need_half_cut=(nuki == 2), need_full_cut=(nuki == 1),
        need_emboss=bool(emb), need_back_print=bool(ura), need_variable=bool(numbering),
        need_haku=(haku_tanka > 0),
    )
    _machine_block(cands, excl, haku_cands)

    with st.expander("ロール仕上・粗利率・送料・外注原価（詳細）"):
        r1 = st.columns(3)
        roll_per = r1[0].number_input("ロール1巻あたり枚数", min_value=1, value=500, step=50, key="h_rp")
        shikan = r1[1].number_input("ロール紙管原価", min_value=0, value=10, step=1, key="h_sk")
        soryo = r1[2].number_input("送料（1箱原価）", min_value=0, value=1000, step=100, key="h_so")
        r2 = st.columns(3)
        mm = r2[0].number_input("材料 粗利率(%)", 0, 99, 30, key="h_mm")
        mo = r2[1].number_input("その他 粗利率(%)", 0, 99, 20, key="h_mo")
        mou = r2[2].number_input("外注・色校正 粗利率(%)", 0, 99, 30, key="h_mou")
        r3 = st.columns(4)
        hanshita = r3[0].number_input("版下 原価", min_value=0, value=0, step=100, key="h_hs")
        han = r3[1].number_input("版代 原価", min_value=0, value=0, step=100, key="h_hn")
        kata = r3[2].number_input("型代 原価", min_value=0, value=0, step=100, key="h_kt")
        seihan = r3[3].number_input("色校正 原価", min_value=0, value=0, step=100, key="h_sh")

    qty = st.number_input("数量（枚数）", min_value=1, value=1000, step=100, key="h_qty")

    base = dict(width=width, height=height, colors=int(colors), machine=machine,
                genshi_tanka=genshi_tanka, pp_tanka=pp_tanka, haku_tanka=haku_tanka,
                numbering=numbering, menzuke_w=int(mw), menzuke_p=int(mp),
                plate_change=int(plate_change), color_change=int(color_change),
                soryo=soryo, nuki=nuki, finish=finish, roll_per=int(roll_per), shikan=shikan,
                margin_material=int(mm), margin_other=int(mo), margin_outsource=int(mou),
                ura_print=ura, half_lami=half, emboss=emb, add_process=int(add_process),
                hanshita_cost=hanshita, han_cost=han, kata_cost=kata, seihan_cost=seihan)

    st.divider()
    st.markdown("#### 計算結果")
    r0 = pr.estimate_hiraatsu(dict(base, unit_price_manual=None), qty=int(qty))
    _cost_block(r0)
    if st.checkbox("参考単価をそのまま使う", value=True, key="h_useref"):
        unit = round(r0["ref_unit"], 2)
    else:
        unit = st.number_input("印刷代 単価（手入力）", min_value=0.0,
                               value=round(r0["ref_unit"], 2), step=0.1, key="h_unit")
    r = pr.estimate_hiraatsu(dict(base, unit_price_manual=unit), qty=int(qty))

    g = st.columns(5)
    g[0].metric("印刷代（提出用）", yen(r["print_submit"]))
    g[1].metric("版下代", yen(r["hanshita"]))
    g[2].metric("版代", yen(r["han"]))
    g[3].metric("型代", yen(r["kata"]))
    g[4].metric("色校正代", yen(r["seihan"]))

    nuki_lbl = "全抜き" if nuki == 1 else "ハーフカット"
    fin_lbl = "シート" if finish == 1 else "ロール"
    spec = f"平圧・間欠 / {int(width)}×{int(height)}mm / {int(colors)}色 / {int(mw)}×{int(mp)}面付 / {nuki_lbl} / {fin_lbl}"
    lines = [("印刷代", r["print_submit"]), ("版下代", r["hanshita"]),
             ("版代", r["han"]), ("型代", r["kata"]), ("色校正代", r["seihan"])]
    return r, spec, unit, lines, int(qty)


def render_masking():
    st.markdown("#### 基本仕様")
    a = st.columns(4)
    width = a[0].number_input("製品の幅 (mm)", min_value=1.0, value=15.0, step=1.0, key="m_w")
    dobu = a[1].number_input("ドブ幅 (mm)", min_value=0.0, value=3.0, step=1.0, key="m_d")
    maki = a[2].number_input("巻きメーター (m)", min_value=1.0, value=5.0, step=1.0, key="m_mk")
    qty = a[3].number_input("数量（巻）", min_value=1, value=1000, step=100, key="m_qty")

    with st.expander("原価単価・営業利益（詳細）"):
        c1 = st.columns(3)
        data_cost = c1[0].number_input("印刷用データ作成", min_value=0, value=3000, step=500, key="m_dc")
        proof = c1[1].number_input("色校正（一律）", min_value=0, value=20000, step=1000, key="m_pf")
        eigyo = c1[2].number_input("営業利益率(%)", 0, 99, 0, key="m_eg")
        masters_ui.material_picker("m_gt")
        c2 = st.columns(4)
        genshi = c2[0].number_input("原紙 平米単価", min_value=0.0, value=49.0, step=1.0, key="m_gt")
        nen = c2[1].number_input("粘着剤 単価", min_value=0.0, value=7.0, step=0.5, key="m_nen")
        nis = c2[2].number_input("シリコンニス 単価", min_value=0.0, value=26.4, step=0.1, key="m_nis")
        ink = c2[3].number_input("インク 単価", min_value=0.0, value=24.0, step=1.0, key="m_ink")
        c3 = st.columns(4)
        pj_print = c3[0].number_input("PJ印刷 単価", min_value=0.0, value=5.8, step=0.1, key="m_pj")
        emul = c3[1].number_input("エマルジョン印刷 単価", min_value=0.0, value=8.0, step=0.5, key="m_em")
        soryo = c3[2].number_input("送料（1箱）", min_value=0, value=1500, step=100, key="m_so")
        pp = c3[3].number_input("PPシート＋内職 単価", min_value=0.0, value=8.0, step=0.5, key="m_pp")

    inp = dict(width=width, dobu=dobu, maki_m=maki, data_cost=data_cost, proof=proof,
               genshi_tanka=genshi, nenchaku=nen, nis=nis, ink=ink, pj_print=pj_print,
               emul_print=emul, soryo=soryo, pp_naishoku=pp, eigyo_rieki=int(eigyo))
    r = pr.estimate_masking(inp, qty=int(qty))

    st.divider()
    st.markdown("#### 計算結果")
    c = st.columns(4)
    c[0].metric("印刷列数", f"{r['retsu']} 列")
    c[1].metric("原紙使用メーター", f"{r['use_m']:,.0f} m")
    c[2].metric("小計（原価積み上げ）", yen(r["subtotal"]))
    c[3].metric("単価（1巻）", f"¥{r['unit']:.2f}")
    with st.expander("原価の内訳を見る"):
        _breakdown_table(r["breakdown"])
    st.caption("🖨️ 印刷機について：マスキングテープでシームレス印刷やシリアル可変が必要な場合は「PJ」機が想定されます。")

    spec = f"マスキングテープ / 幅{int(width)}mm × {maki:g}m巻"
    lines = [("印刷代", r["total"])]
    return r, spec, r["unit"], lines, int(qty)


def render_konica():
    st.markdown("#### 基本仕様")
    a = st.columns(4)
    width = a[0].number_input("寸法ヨコ W (mm)", min_value=0.0, value=50.0, step=1.0, key="k_w")
    height = a[1].number_input("寸法タテ P (mm)", min_value=0.0, value=30.0, step=1.0, key="k_h")
    types = a[2].number_input("種類（同一型・同一仕上）", min_value=1, value=1, step=1, key="k_ty")
    version = _sel("版", {"新版": 1, "改版": 2, "再版": 3}, "k_ver")
    b = st.columns(4)
    print_w = b[0].number_input("印刷面付ヨコ", min_value=1, value=1, step=1, key="k_pw")
    nuki_w = b[1].number_input("抜き面付ヨコ", min_value=1, value=1, step=1, key="k_nw")
    nuki_p = b[2].number_input("抜き面付タテ", min_value=1, value=1, step=1, key="k_np")
    nuki = _sel("抜き", {"ハーフカット": 2, "全抜き": 1}, "k_nuki")

    st.markdown("#### 材料（平米単価）・仕上")
    masters_ui.material_picker("k_gt")
    m = st.columns(4)
    genshi_tanka = m[0].number_input("原紙 平米単価", min_value=0.0, step=1.0, key="k_gt")
    pp_tanka = m[1].number_input("PP 平米単価", min_value=0.0, value=0.0, step=1.0, key="k_pt")
    finish = _sel("仕上形状", {"シート": 1, "ロール": 2}, "k_fin")
    slitter_hon = m[3].number_input("スリッター本数", min_value=1, value=1, step=1, key="k_sl")

    cands, excl, _ = machine_select.suggest_machines(
        colors=4, width=width, height=height, need_half_cut=(nuki == 2), need_full_cut=(nuki == 1),
    )
    _machine_block(cands, excl, [], restrict_category="オンデマンド")
    st.caption("※ コニカミノルタ方式はオンデマンド機（コニカ／PJ）が対象です。シリアル可変・シームレス印刷が必要な場合はPJをご検討ください。")

    with st.expander("ロール仕上・粗利率・送料・外注原価（詳細）"):
        r1 = st.columns(4)
        roll_per = r1[0].number_input("ロール1巻あたり枚数", min_value=1, value=500, step=50, key="k_rp")
        shikan = r1[1].number_input("ロール紙管原価", min_value=0, value=10, step=1, key="k_sk")
        soryo = r1[2].number_input("送料（1箱原価）", min_value=0, value=1000, step=100, key="k_so")
        add_process = r1[3].number_input("追加工程数", min_value=0, value=0, step=1, key="k_ap")
        r2 = st.columns(3)
        mm = r2[0].number_input("材料 粗利率(%)", 0, 99, 30, key="k_mm")
        mo = r2[1].number_input("その他 粗利率(%)", 0, 99, 20, key="k_mo")
        mou = r2[2].number_input("外注・色校正 粗利率(%)", 0, 99, 30, key="k_mou")
        r3 = st.columns(3)
        hanshita = r3[0].number_input("版下 原価", min_value=0, value=0, step=100, key="k_hs")
        kata = r3[1].number_input("型代 原価", min_value=0, value=0, step=100, key="k_kt")
        seihan = r3[2].number_input("色校正 原価", min_value=0, value=0, step=100, key="k_sh")

    qty = st.number_input("数量（枚数）", min_value=1, value=1000, step=100, key="k_qty")

    base = dict(width=width, height=height, types=int(types), genshi_tanka=genshi_tanka,
                pp_tanka=pp_tanka, version=version, print_w=int(print_w), nuki_w=int(nuki_w),
                nuki_p=int(nuki_p), slitter_hon=int(slitter_hon), soryo=soryo, nuki=nuki,
                finish=finish, roll_per=int(roll_per), shikan=shikan, add_process=int(add_process),
                margin_material=int(mm), margin_other=int(mo), margin_outsource=int(mou),
                hanshita_cost=hanshita, kata_cost=kata, seihan_cost=seihan)

    st.divider()
    st.markdown("#### 計算結果")
    r0 = pr.estimate_konica(dict(base, unit_price_manual=None), qty=int(qty))
    _cost_block(r0)
    if st.checkbox("参考単価をそのまま使う", value=True, key="k_useref"):
        unit = round(r0["ref_unit"], 2)
    else:
        unit = st.number_input("印刷代 単価（手入力）", min_value=0.0,
                               value=round(r0["ref_unit"], 2), step=0.1, key="k_unit")
    r = pr.estimate_konica(dict(base, unit_price_manual=unit), qty=int(qty))

    g = st.columns(4)
    g[0].metric("印刷代（提出用）", yen(r["print_submit"]))
    g[1].metric("版下代", yen(r["hanshita"]))
    g[2].metric("型代", yen(r["kata"]))
    g[3].metric("色校正代", yen(r["seihan"]))

    nuki_lbl = "全抜き" if nuki == 1 else "ハーフカット"
    fin_lbl = "シート" if finish == 1 else "ロール"
    spec = f"コニカミノルタ / {int(width)}×{int(height)}mm / 印刷面付{int(print_w)} / {nuki_lbl} / {fin_lbl}"
    lines = [("印刷代", r["print_submit"]), ("版下代", r["hanshita"]),
             ("型代", r["kata"]), ("色校正代", r["seihan"])]
    return r, spec, unit, lines, int(qty)


# =====================================================================
# 見積書（テキスト／Excel）
# =====================================================================
def _quote_no(seq=1):
    d = datetime.now()
    return f"SD-{d.strftime('%Y%m%d')}-{seq:03d}"


def build_quote_text(name, spec, qty, unit, lines, meta):
    pretax = sum(a for _, a in lines)
    tax = round(pretax * TAX_RATE)
    total = pretax + tax
    L = []
    if meta.get("customer"):
        L.append(meta["customer"] + " 御中")
    L += ["", "お見積りのご回答を申し上げます。", ""]
    if meta.get("subject"):
        L.append("件名：" + meta["subject"])
    L += [f"見積番号：{meta.get('quote_no','')}", ""]
    L.append(f"■ {name}")
    L.append(f"　仕様：{spec}")
    L.append(f"　数量：{qty:,}／単価：¥{unit:.2f}")
    L.append("────────────────")
    for lbl, amt in lines:
        if amt:
            L.append(f"{lbl}：{yen(amt)}")
    L.append(f"小計（税抜）：{yen(pretax)}")
    L.append(f"消費税（10%）：{yen(tax)}")
    L.append(f"合計金額（税込）：{yen(total)}")
    L += ["────────────────", ""]
    if meta.get("remarks"):
        L.append("備考：" + meta["remarks"])
    L.append("※金額は実際の見積計算フォーマットに基づく算定です。")
    L += ["", ISSUER["company"] + (("　" + meta["staff"]) if meta.get("staff") else ""), ISSUER["tel"]]
    return "\n".join(L)


def build_quote_excel(name, spec, qty, unit, lines, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    pretax = sum(a for _, a in lines)
    tax = round(pretax * TAX_RATE)
    total = pretax + tax
    wb = Workbook(); ws = wb.active; ws.title = "見積書"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([26, 40, 12, 16], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    thin = Side(style="thin", color="AAAAAA"); bd = Border(thin, thin, thin, thin)
    issue = meta.get("issue_date") or date.today()
    until = issue + timedelta(days=int(meta.get("valid_days", 30)))
    ws["A1"] = "御 見 積 書"; ws["A1"].font = Font(bold=True, size=22)
    ws.merge_cells("A1:C1")
    ws["D1"] = f"見積番号：{meta.get('quote_no','')}"; ws["D1"].alignment = Alignment(horizontal="right")
    ws["D2"] = f"発行日：{issue:%Y年%m月%d日}"; ws["D2"].alignment = Alignment(horizontal="right")
    ws["D3"] = f"有効期限：{until:%Y年%m月%d日}"; ws["D3"].alignment = Alignment(horizontal="right")
    cust = (meta.get("customer") or "").strip()
    ws["A4"] = (cust + "　御中") if cust else "　御中"
    ws["A4"].font = Font(bold=True, size=14, underline="single")
    if meta.get("subject"):
        ws["A5"] = "件名：" + meta["subject"]
    ws["A6"] = "下記の通りお見積り申し上げます。"
    ws["A7"] = "お見積金額（税込）"; ws["A7"].font = Font(bold=True, size=12)
    ws["C7"] = total; ws["C7"].number_format = '"¥"#,##0'; ws["C7"].font = Font(bold=True, size=16)
    ws.merge_cells("C7:D7")
    row = 9
    hdr = ["品名／項目", "仕様", "数量", "金額"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = PatternFill("solid", fgColor="1F3B57"); c.font = Font(color="FFFFFF", bold=True)
        c.border = bd; c.alignment = Alignment(horizontal="center")
    row += 1
    ws.cell(row=row, column=1, value=name).border = bd
    ws.cell(row=row, column=2, value=spec).border = bd
    q = ws.cell(row=row, column=3, value=qty); q.number_format = "#,##0"; q.border = bd
    a = ws.cell(row=row, column=4, value=round(lines[0][1])); a.number_format = '"¥"#,##0'; a.border = bd
    row += 1
    for lbl, amt in lines[1:]:
        if not amt:
            continue
        ws.cell(row=row, column=1, value=lbl).border = bd
        ws.cell(row=row, column=2, value="").border = bd
        ws.cell(row=row, column=3, value="").border = bd
        c = ws.cell(row=row, column=4, value=round(amt)); c.number_format = '"¥"#,##0'; c.border = bd
        row += 1
    for lbl, val, bold in [("小計（税抜）", pretax, False), ("消費税(10%)", tax, False), ("合計（税込）", total, True)]:
        lc = ws.cell(row=row, column=3, value=lbl); lc.alignment = Alignment(horizontal="right")
        vc = ws.cell(row=row, column=4, value=round(val)); vc.number_format = '"¥"#,##0'
        if bold:
            lc.font = Font(bold=True); vc.font = Font(bold=True)
        lc.border = bd; vc.border = bd
        row += 1
    row += 1
    if meta.get("remarks"):
        ws.cell(row=row, column=1, value="備考：" + meta["remarks"]); row += 1
    row += 1
    ws.cell(row=row, column=1, value=ISSUER["company"]).font = Font(bold=True, size=12); row += 1
    ws.cell(row=row, column=1, value=f"{ISSUER['postal']} {ISSUER['address']}"); row += 1
    ws.cell(row=row, column=1, value=ISSUER["tel"]); row += 1
    if meta.get("staff"):
        ws.cell(row=row, column=1, value="担当：" + meta["staff"])
    bio = BytesIO(); wb.save(bio); bio.seek(0)
    return bio, total


def render_quote_output(method, name, spec, qty, unit, lines):
    st.divider()
    st.markdown("#### 見積書の情報")
    a = st.columns(4)
    customer = a[0].text_input("宛先", placeholder="株式会社サンプル", key="rq_cust")
    subject = a[1].text_input("件名", placeholder="ラベル印刷 御見積", key="rq_subj")
    staff = a[2].text_input("担当者", placeholder="深見", key="rq_staff")
    valid_days = a[3].number_input("有効期限(日数)", 7, 180, 30, key="rq_valid")
    remarks = st.text_input("備考（見積書に載る）", placeholder="完全データ支給／一括納品想定 など", key="rq_rem")

    pretax = sum(x for _, x in lines)
    tax = round(pretax * TAX_RATE)
    st.markdown(
        f'<div style="background:#EDF1DC;border:1px solid #B7C77E;border-radius:10px;'
        f'padding:12px 16px;display:flex;justify-content:space-between;align-items:center">'
        f'<b style="color:#3B4552">合計金額（税込）</b>'
        f'<span style="font-size:24px;font-weight:800;color:#1E2A36">{yen(pretax + tax)}</span></div>',
        unsafe_allow_html=True)
    st.caption(f"内訳：小計（税抜） {yen(pretax)}／消費税(10%) {yen(tax)}")

    meta = dict(quote_no=_quote_no(), issue_date=date.today(), valid_days=int(valid_days),
                customer=customer, subject=subject, staff=staff, remarks=remarks)

    out = st.columns(2)
    with out[0]:
        xlsx, _ = build_quote_excel(name, spec, qty, unit, lines, meta)
        st.download_button("📄 見積書をExcelで出力", data=xlsx,
                           file_name=f"見積書_{customer or 'お客様'}_{meta['quote_no']}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           type="primary", width="stretch")
    with out[1]:
        if st.button("📋 メール返信用テキストを作る", width="stretch"):
            st.session_state["rq_show_text"] = True
    if st.session_state.get("rq_show_text"):
        st.text_area("見積回答テキスト（コピーして使えます）",
                     value=build_quote_text(name, spec, qty, unit, lines, meta), height=280)

    if st.button("💾 この見積を履歴に保存", width="stretch"):
        store.add_history({
            "quote_no": meta["quote_no"], "date": datetime.now().strftime("%Y/%m/%d"),
            "customer": customer or "", "name": name, "method": method,
            "qty": int(qty), "unit": round(unit, 2), "total": pretax + tax, "status": "下書き"})
        st.success(f"履歴に保存しました（{meta['quote_no']}）。")


# =====================================================================
def render():
    st.markdown("### 見積作成")
    st.markdown("**① まず入力方法を選んでください**")
    _PASTE, _MANUAL, _PDF = "📋 問い合わせ文を貼り付け", "✏️ 仕様を直接入力", "📄 図面・注文書PDFから"
    mode = st.radio("入力方法", [_PASTE, _MANUAL, _PDF], horizontal=True, key="rq_mode",
                    label_visibility="collapsed",
                    captions=["メール等の問い合わせ文を貼って自動反映", "サイズ・数量などを直接入力",
                              "図面・注文書PDFを読み取って自動入力"])
    st.divider()
    with st.container(border=True):
        if mode == _PASTE:
            render_paste_flow()
        elif mode == _PDF:
            pdf_intake.render_intake()
        else:
            st.markdown("#### ✏️ 仕様を直接入力")
            st.caption("下の「計算方式」を選び、寸法・数量などを入力してください。")
        _template_bar()

    st.divider()
    st.markdown("**② 計算方式と仕様**")
    top = st.columns([2, 2])
    method = top[0].selectbox("計算方式",
                              ["平圧・間欠印刷", "マスキングテープ", "コニカミノルタ オンデマンド"],
                              key="rq_method")
    name = top[1].text_input("品名", value="シール・ラベル", key="rq_name")
    with st.container(border=True):
        if method == "平圧・間欠印刷":
            r, spec, unit, lines, qty = render_hiraatsu()
        elif method == "マスキングテープ":
            r, spec, unit, lines, qty = render_masking()
        else:
            r, spec, unit, lines, qty = render_konica()
        render_quote_output(method, name or "品目", spec, qty, unit, lines)
    st.caption("※ この計算は実際のExcel見積フォーマットと1円単位まで一致することを検証済みです。")
