# -*- coding: utf-8 -*-
"""
Excel見積書の生成（openpyxl）

社員モードで積み上げた複数品目と、宛先・件名・値引きなどから、
そのままお客様に出せる体裁の見積書(.xlsx)を組み立てて BytesIO で返す。

発行元（シール堂印刷）の情報は下の ISSUER で管理。ロゴや押印は後フェーズ。
金額は pricing.py で計算済みの数値をそのまま流し込むだけ（ここでは再計算しない）。
"""

from io import BytesIO
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# 発行元情報（実際の連絡先はサイトの会社概要より）
ISSUER = {
    "company": "株式会社シール堂印刷",
    "postal": "〒141-0031",
    "address": "東京都品川区西五反田4-27-10 印刷産業ビル",
    "tel": "TEL 03-3493-2131",
    "fax": "",
    "note": "シール・ラベル・ステッカーの企画製造販売",
}

# 色・罫線の共通設定
_HEADER_FILL = PatternFill("solid", fgColor="1F3B57")   # 濃紺
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_TITLE_FONT = Font(bold=True, size=22)
_TOTAL_FILL = PatternFill("solid", fgColor="EAF0F6")
_thin = Side(style="thin", color="AAAAAA")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_YEN = '"¥"#,##0'


def _yen(cell):
    cell.number_format = _YEN
    cell.alignment = Alignment(horizontal="right")


def build_quote_excel(items, totals, meta):
    """
    items  : estimate() の返り値（valid な品目）のリスト
    totals : apply_discount() の返り値（合計・値引き・税など）
    meta   : dict
        quote_no       見積番号
        issue_date     発行日(date) ※省略時は今日
        valid_days     有効期限の日数（既定30日）
        customer       宛先（会社名・お名前）
        subject        件名
        staff          担当者名
        remarks        備考（自由記入）
        show_profit    True なら社内用に原価・粗利の欄を付ける
    戻り値 : BytesIO（.xlsx バイナリ）
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "見積書"
    ws.sheet_view.showGridLines = False

    issue = meta.get("issue_date") or date.today()
    valid_days = int(meta.get("valid_days", 30))
    valid_until = issue + timedelta(days=valid_days)
    show_profit = bool(meta.get("show_profit"))

    # 列幅（品名/仕様/数量/単価/金額 (+原価/粗利)）
    widths = [26, 40, 10, 14, 16]
    if show_profit:
        widths += [14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last_col = len(widths)
    last_col_letter = get_column_letter(last_col)

    r = 1
    # タイトル（右端の1列は見積番号などの情報欄に空けておく＝結合範囲を重ねない）
    ws.merge_cells(f"A{r}:{get_column_letter(max(1, last_col - 1))}{r}")
    c = ws[f"A{r}"]
    c.value = "御 見 積 書"
    c.font = _TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r].height = 34

    # 右上（常に右端列）に見積番号・発行日・有効期限
    info_col = last_col_letter
    ws[f"{info_col}{r}"] = f"見積番号：{meta.get('quote_no', '')}"
    ws[f"{info_col}{r}"].alignment = Alignment(horizontal="right")
    ws[f"{info_col}{r+1}"] = f"発行日：{issue.strftime('%Y年%m月%d日')}"
    ws[f"{info_col}{r+1}"].alignment = Alignment(horizontal="right")
    ws[f"{info_col}{r+2}"] = f"有効期限：{valid_until.strftime('%Y年%m月%d日')}"
    ws[f"{info_col}{r+2}"].alignment = Alignment(horizontal="right")
    r += 3

    # 宛先
    r += 1
    customer = (meta.get("customer") or "").strip()
    honorific = "御中" if customer and not customer.endswith(("様", "御中")) else ""
    ws[f"A{r}"] = f"{customer}　{honorific}".strip() if customer else "　御中"
    ws[f"A{r}"].font = Font(bold=True, size=14, underline="single")
    r += 2

    # 件名・リード文
    if meta.get("subject"):
        ws[f"A{r}"] = f"件名：{meta['subject']}"
        ws[f"A{r}"].font = Font(size=11)
        r += 1
    ws[f"A{r}"] = "下記の通りお見積り申し上げます。"
    ws[f"A{r}"].font = Font(size=11)
    r += 2

    # 合計金額（税込）を目立たせる帯
    ws.merge_cells(f"A{r}:B{r}")
    ws[f"A{r}"] = "お見積金額（税込）"
    ws[f"A{r}"].font = Font(bold=True, size=12)
    ws[f"A{r}"].fill = _TOTAL_FILL
    ws.merge_cells(f"C{r}:{last_col_letter}{r}")
    tot = ws[f"C{r}"]
    tot.value = totals["total"]
    tot.number_format = _YEN
    tot.font = Font(bold=True, size=16)
    tot.alignment = Alignment(horizontal="right", vertical="center")
    tot.fill = _TOTAL_FILL
    ws.row_dimensions[r].height = 26
    for col in range(1, last_col + 1):
        ws.cell(row=r, column=col).border = _BORDER
    r += 2

    # 明細ヘッダー
    headers = ["品名", "仕様", "数量", "単価", "金額"]
    if show_profit:
        headers += ["原価", "粗利"]
    header_row = r
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=i, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    ws.row_dimensions[r].height = 20
    r += 1

    # 明細行
    for idx, it in enumerate(items, start=1):
        name = it.get("name") or f"シール・ラベル {idx}"
        ws.cell(row=r, column=1, value=name).border = _BORDER
        spec_cell = ws.cell(row=r, column=2, value=it.get("spec_text", ""))
        spec_cell.border = _BORDER
        spec_cell.alignment = Alignment(wrap_text=True, vertical="top")
        qcell = ws.cell(row=r, column=3, value=it["qty"])
        qcell.number_format = "#,##0\"枚\""
        qcell.alignment = Alignment(horizontal="right")
        qcell.border = _BORDER
        ucell = ws.cell(row=r, column=4, value=round(it["unit_price"], 2))
        ucell.number_format = '"¥"#,##0.00'
        ucell.alignment = Alignment(horizontal="right")
        ucell.border = _BORDER
        acell = ws.cell(row=r, column=5, value=it["subtotal"])
        _yen(acell)
        acell.border = _BORDER
        if show_profit:
            cc = ws.cell(row=r, column=6, value=it["cost"])
            _yen(cc); cc.border = _BORDER
            pc = ws.cell(row=r, column=7, value=it["profit"])
            _yen(pc); pc.border = _BORDER
        r += 1

    # 合計ゾーン（右寄せ）。ラベルは単価列、値は金額列に置く
    label_col = 4
    value_col = 5

    def _line(label, value, bold=False, money=True):
        nonlocal r
        lc = ws.cell(row=r, column=label_col, value=label)
        lc.alignment = Alignment(horizontal="right")
        if bold:
            lc.font = Font(bold=True)
        vc = ws.cell(row=r, column=value_col, value=value)
        if money:
            _yen(vc)
        else:
            vc.alignment = Alignment(horizontal="right")
        if bold:
            vc.font = Font(bold=True)
        vc.border = _BORDER
        lc.border = _BORDER
        r += 1

    _line("小計（税抜）", totals["items_subtotal"])
    if totals["discount"]:
        _line("値引き", -totals["discount"])
        _line("値引後小計", totals["subtotal_after"])
    _line(f"消費税（10%）", totals["tax"])
    _line("合計（税込）", totals["total"], bold=True)

    # 社内用：粗利まとめ
    if show_profit:
        r += 1
        ws.cell(row=r, column=1, value="【社内用】原価・粗利まとめ").font = Font(bold=True, color="B00020")
        r += 1
        rate_pct = f"{totals['profit_rate']*100:.1f}%"
        ws.cell(row=r, column=1, value=f"原価合計：¥{totals['cost']:,}／粗利：¥{totals['profit']:,}（粗利率 {rate_pct}）")
        if totals["below_cost"]:
            r += 1
            warn = ws.cell(row=r, column=1, value="⚠ 原価割れ（赤字）です。値引きを見直してください。")
            warn.font = Font(bold=True, color="B00020")
        r += 1

    # 備考
    r += 1
    ws.cell(row=r, column=1, value="備考").font = Font(bold=True)
    r += 1
    default_remark = "※本見積書の金額は仮の料金設定に基づく参考値です。正式なご発注前に最終金額をご確認ください。"
    remarks = meta.get("remarks") or ""
    remark_text = (remarks + ("\n" if remarks else "") + default_remark).strip()
    ws.merge_cells(f"A{r}:{last_col_letter}{r+2}")
    rc = ws[f"A{r}"]
    rc.value = remark_text
    rc.alignment = Alignment(wrap_text=True, vertical="top")
    rc.border = _BORDER
    r += 4

    # 発行元
    ws.cell(row=r, column=1, value=ISSUER["company"]).font = Font(bold=True, size=12)
    r += 1
    for line in [f"{ISSUER['postal']} {ISSUER['address']}", ISSUER["tel"], ISSUER["note"]]:
        if line.strip():
            ws.cell(row=r, column=1, value=line)
            r += 1
    if meta.get("staff"):
        ws.cell(row=r, column=1, value=f"担当：{meta['staff']}")

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
